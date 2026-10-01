#!/usr/bin/env python3
"""
Tabela única do artigo: receita per capita municipal do Paraná (ISS + cota-parte
do ICMS em 2025, substituídos pelo IBS depois) nos municípios de maior e menor
valor em 2025, com colunas 2025, 2029, 2033, 2050 e 2077, e medidas de
dispersão (razão máx./mín. e Gini) calculadas sobre todos os municípios elegíveis.

Reaproveita montar() de build-tabelas-artigo-pr.py (base final auditada, ISS com
correções do PIT/TCE-PR, R$ constantes de 2025, população fixa). 2029-2033 vêm
direto da base final auditada; 2050 e 2077 seguem a mesma fórmula de 2033 com
ibs-projecao-longo-prazo.json (Seguro-Receita: fatia de 2033 x fundo municipal
do PR no ano).

Uso: python3 build-tabela-unica-pr.py
"""
import csv, importlib.util, json, statistics
from pathlib import Path

HERE = Path(__file__).parent
spec = importlib.util.spec_from_file_location("bt", HERE / "build-tabelas-artigo-pr.py")
bt = importlib.util.module_from_spec(spec); spec.loader.exec_module(bt)

ANOS = [2025, 2029, 2033, 2050, 2077]
N = 10
OUT = bt.OUT


def dma_rel(v):
    m = statistics.mean(v)
    return sum(abs(x - m) for x in v) / len(v) / m


def medidas(v):
    return {"Razão máx./mín.": (max(v) / min(v), 1), "Índice de Gini": (bt.gini(v), 3)}


def main():
    linhas, eleg = bt.montar()
    final, _, _, _, lp, seg = bt.carregar()
    soma_seg = sum(float(r["seguro_2033_final"]) for r in final.values())
    for l in linhas:
        r = final[l["codigo_ibge"]]
        rec = {2025: l["base_2025"], 2029: float(r["receita_2029_final"]),
               2033: float(r["receita_2033_final"])}
        cpt, dest = float(r["cpt_final_pct"]) / 100, float(r["destino_pct"]) / 100
        fatia = float(r["seguro_2033_final"]) / soma_seg
        for y in (2050, 2077):
            p = lp[y]
            fundo = seg[str(y)]["repasse_por_uf"]["PR"]["municipio"]
            rec[y] = ((1 - p["ca"]) * p["ibs_historico"] * cpt + p["ibs_destino_liquido"] * dest
                      + (1 - p["ca"]) * fatia * fundo)
        for y in ANOS:
            l[f"pc{y}"] = rec[y] / l["pop"]
    assert all(abs(l["pc2033"] - l["pc_2033"]) < 1e-6 and abs(l["pc2077"] - l["pc_2077"]) < 1e-6 for l in linhas)

    rank = sorted(eleg, key=lambda l: -l["pc2025"])
    ext = rank[:N] + rank[-N:]
    disp = {y: medidas([l[f"pc{y}"] for l in eleg]) for y in ANOS}
    json.dump({"n": len(eleg), "extremos": [{k: l[k] for k in ["municipio", "pop"] + [f"pc{y}" for y in ANOS]} for l in ext],
               "dispersao": {str(y): {k: v[0] for k, v in d.items()} for y, d in disp.items()}},
              open(OUT / "resumo-tabela-unica.json", "w"), ensure_ascii=False, indent=1)
    with open(OUT / "municipios-percapita-pr-anos.csv", "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        ids = {l["codigo_ibge"] for l in eleg}
        w.writerow(["codigo_ibge", "municipio", "pop_media"] + [f"receita_pc_{y}" for y in ANOS] + ["elegivel"])
        for l in sorted(linhas, key=lambda l: -l["pc2025"]):
            w.writerow([l["codigo_ibge"], l["municipio"], round(l["pop"], 1)] + [round(l[f"pc{y}"], 2) for y in ANOS]
                       + [l["codigo_ibge"] in ids])
    gerar_docx(ext, disp, len(eleg))
    return ext, disp, len(eleg)


def gerar_docx(ext, disp, n_eleg):
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, Cm
    br = bt._br
    doc = Document()
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21.0), Cm(29.7)
    s.left_margin = s.right_margin = Cm(2.0)
    st = doc.styles["Normal"]; st.font.name = "Times New Roman"; st.font.size = Pt(10)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

    def par(t, bold=False, italic=False, size=10, after=4, align=None):
        p = doc.add_paragraph(); r = p.add_run(t); r.bold, r.italic = bold, italic; r.font.size = Pt(size)
        p.paragraph_format.space_after = Pt(after)
        if align: p.alignment = align
        return p

    def borda(c, **kw):
        tcPr = c._tc.get_or_add_tcPr()
        b = tcPr.find(qn("w:tcBorders"))
        if b is None:
            b = OxmlElement("w:tcBorders"); tcPr.append(b)
        for lado, (v, sz) in kw.items():
            e = OxmlElement(f"w:{lado}"); e.set(qn("w:val"), v); e.set(qn("w:sz"), str(sz)); b.append(e)

    def sombra(c, cor="E7E6E6"):
        e = OxmlElement("w:shd"); e.set(qn("w:val"), "clear"); e.set(qn("w:color"), "auto"); e.set(qn("w:fill"), cor)
        c._tc.get_or_add_tcPr().append(e)

    def celula(c, txt, bold=False, italic=False, align="r", size=9):
        c.text = ""; p = c.paragraphs[0]; r = p.add_run(txt); r.bold, r.italic = bold, italic; r.font.size = Pt(size)
        p.alignment = {"l": WD_ALIGN_PARAGRAPH.LEFT, "r": WD_ALIGN_PARAGRAPH.RIGHT, "c": WD_ALIGN_PARAGRAPH.CENTER}[align]
        p.paragraph_format.space_after = Pt(0); p.paragraph_format.space_before = Pt(1)

    par("Tabela 1 – Receita per capita de ISS e cota-parte do ICMS (2025) e do IBS municipal (2029-2077) nos municípios "
        "paranaenses de maior e menor valor em 2025, com medidas de dispersão (R$ de 2025)", bold=True, size=10, after=6)

    larg = [4.3, 1.9] + [1.75] * len(ANOS) + [2.05]
    ncol = len(larg)
    t = doc.add_table(rows=0, cols=ncol); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    tblPr = t._tbl.tblPr
    lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblPr.append(lay)
    def nova():
        r = t.add_row()
        for i, w in enumerate(larg): r.cells[i].width = Cm(w)
        return r.cells
    def var(a, b):
        v = f"{(b / a - 1) * 100:+,.1f}%".replace(",", "X").replace(".", ",").replace("X", ".")
        return v.replace("-", "\u2212")

    # cabeçalho em dois níveis (células de rótulo mescladas na vertical)
    c = nova()
    g = c[2].merge(c[2 + len(ANOS) - 1]); celula(g, "Receita per capita (R$)", bold=True, align="c")
    borda(g, bottom=("single", 4))
    c2 = nova()
    for i in range(2, 2 + len(ANOS)): celula(c2[i], str(ANOS[i - 2]), bold=True, align="c")
    for i, h in ((0, "Município"), (1, "População"), (ncol - 1, "Variação\n2025-2077")):
        m = c[i].merge(c2[i]); celula(m, h, bold=True, align="l" if i == 0 else "c")
    for cel in c: borda(cel, top=("single", 12))
    for cel in c2: borda(cel, bottom=("single", 6))

    n = N
    def linha_mun(rot, l):
        c = nova()
        celula(c[0], rot, align="l"); celula(c[1], br(l["pop"]))
        for k, y in enumerate(ANOS): celula(c[2 + k], br(l[f"pc{y}"]))
        celula(c[-1], var(l["pc2025"], l["pc2077"]))
    for i, l in enumerate(ext[:n]): linha_mun(f"{i + 1}. {l['municipio']}", l)
    c = nova()
    celula(c[0], f"... ({n_eleg - 2 * n} municípios omitidos)", italic=True, align="l")
    for k in range(1, ncol): celula(c[k], "...", italic=True, align="c")
    for i, l in enumerate(ext[n:]): linha_mun(f"{n_eleg - n + i + 1}. {l['municipio']}", l)

    # linhas de dispersão, destacadas
    for k, nome in enumerate(disp[2025]):
        casas = disp[2025][nome][1]
        c = nova()
        rot = f"{nome} – {n_eleg} municípios"
        celula(c[0], rot, bold=True, align="l", size=9)
        g = c[0].merge(c[1]); celula(g, rot, bold=True, align="l")
        for kk, y in enumerate(ANOS): celula(c[2 + kk], br(disp[y][nome][0], casas), bold=True)
        celula(c[-1], var(disp[2025][nome][0], disp[2077][nome][0]), bold=True)
        for cel in c:
            sombra(cel)
            if k == 0: borda(cel, top=("single", 8))
            if k == len(disp[2025]) - 1: borda(cel, bottom=("single", 12))
    par("", after=2, size=4)
    par("Fonte: elaboração própria, com base na DCA/Siconfi, no Portal da Transparência do Paraná e no PIT/TCE-PR (base municipal final "
        "auditada). Nota: receita per capita = ISS + cota-parte do ICMS em 2025, substituídos pelo IBS municipal de 2029 em diante, conforme o "
        "cronograma do art. 131 do ADCT (EC 132/2023) e PIB real de 2,2% a.a. após 2033; valores em R$ constantes de 2025 e população fixa "
        "(média 2019-2026). Razão máx./mín. = maior valor per capita dividido pelo menor; Índice de Gini (0 = igualdade perfeita). As duas "
        "medidas consideram todos os municípios elegíveis, e não apenas os exibidos; os municípios intermediários foram omitidos (...).",
        italic=False, size=8, after=0)
    doc.save(OUT / "tabela-unica-dispersao-percapita-pr.docx")


if __name__ == "__main__":
    ext, disp, n = main()
    for i, l in enumerate(ext): print(f"{l['municipio']:<24}", *[f"{l[f'pc{y}']:>8.0f}" for y in ANOS])
    for k in disp[2025]: print(f"{k:<36}", *[f"{disp[y][k][0]:>9.3f}" for y in ANOS])
    print(n)
