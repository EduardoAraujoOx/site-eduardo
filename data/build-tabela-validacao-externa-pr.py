#!/usr/bin/env python3
"""
Tabela de validação externa das projeções estaduais: nosso estudo x projeção
do IPEA/Gobetti (R$ bilhões, preços de 2025; base 2025 e 2033), com o Paraná em
destaque, no padrão ABNT (Word).

Três colunas em 2033 permitem separar o efeito do coeficiente do efeito do
método: (i) nosso modelo com o coeficiente de destino próprio (POF/Censo),
(ii) o mesmo modelo com o coeficiente de Gobetti e Monteiro (2023), e (iii) a
projeção do IPEA/Gobetti. Se (ii) ~ (iii), a diferença entre (i) e (iii) vem do
coeficiente, não da mecânica do modelo.

As projeções do IPEA/Gobetti NÃO estão no repositório (repositório público; a
nota técnica é preliminar). O script as lê em tempo de execução da Tabela 12 da
nota técnica do Comsefaz (docx), cujo caminho vai no 1.º argumento.

Uso: python3 build-tabela-validacao-externa-pr.py <nota_tecnica.docx> [saida.docx]
"""
import json
import statistics as st
import sys
from pathlib import Path
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm

HERE = Path(__file__).parent
ROOT = HERE.parent
DESTAQUE = "PR"
LIMIAR_BASE = 5.0   # marca UFs com |Δ base 2025| >= 5%: divergência vem da base, não do coeficiente
NOMES = {'Distrito Federal': 'DF', 'Acre': 'AC', 'Alagoas': 'AL', 'Amazonas': 'AM', 'Amapá': 'AP', 'Bahia': 'BA', 'Ceará': 'CE',
         'Espírito Santo': 'ES', 'Goiás': 'GO', 'Maranhão': 'MA', 'Minas Gerais': 'MG', 'Mato Grosso do Sul': 'MS',
         'Mato Grosso': 'MT', 'Pará': 'PA', 'Paraíba': 'PB', 'Pernambuco': 'PE', 'Piauí': 'PI', 'Paraná': 'PR',
         'Rio de Janeiro': 'RJ', 'Rio Grande do Norte': 'RN', 'Rondônia': 'RO', 'Roraima': 'RR', 'Rio Grande do Sul': 'RS',
         'Santa Catarina': 'SC', 'Sergipe': 'SE', 'São Paulo': 'SP', 'Tocantins': 'TO'}


def le_ipea(caminho):
    """Tabela 12 da nota técnica: linhas por estado, colunas 2025, 2029..2033 em R$ mil."""
    doc = Document(caminho)
    alvo = None
    for t in doc.tables:
        cab = [c.text.strip() for c in t.rows[0].cells]
        if cab[:3] == ["Ano", "2025", "2029"]:
            alvo = t
            break
    assert alvo is not None, "Tabela 12 não encontrada"
    out = {}
    for r in alvo.rows:
        c = [x.text.strip() for x in r.cells]
        if c[0] in NOMES:
            out[NOMES[c[0]]] = [float(x.replace(".", "")) / 1e6 for x in c[1:7]]   # R$ bi: 2025, 2029..2033
    assert len(out) == 27, len(out)
    return out


def monta(ipea):
    R = json.load(open(ROOT / "painelufir/data/resultados-consolidados-ibs.json"))["por_uf_estado"]
    G = json.load(open(HERE / "faixa-phi-dest-estados.json"))["por_uf"]
    linhas = []
    for uf, v in R.items():
        con = v["contrafactual_por_ano"]["2033"]
        def proj(m):   # variacao_por_ano em %, já com o Seguro-Receita refeito para cada método
            return con * (1 + G[uf]["metodos"][m]["variacao_por_ano"]["2033"] / 100) / 1e9
        own = v["pos_por_ano"]["2033"] / 1e9
        assert abs(proj("pof_censo_bruto") - own) / own < 2e-3, uf     # a faixa reproduz o publicado
        i = ipea[uf]
        linhas.append({"uf": uf, "pre": v["pre_2025"] / 1e9, "ipea25": i[0], "own": own, "gob": proj("gobetti_2023"), "ipea33": i[5]})
    for l in linhas:
        l["d25"] = 100 * (l["ipea25"] / l["pre"] - 1)
        l["d_own"] = 100 * (l["ipea33"] / l["own"] - 1)
        l["d_gob"] = 100 * (l["ipea33"] / l["gob"] - 1)
        l["flag"] = abs(l["d25"]) >= LIMIAR_BASE
    return sorted(linhas, key=lambda l: -l["d25"])


def br(x, c=1):
    return f"{x:,.{c}f}".replace(",", "X").replace(".", ",").replace("X", ".")


def sg(x):
    s = f"{x:+.1f}%".replace(".", ",").replace("-", "−")
    return "0,0%" if s in ("+0,0%", "−0,0%") else s


def resumo(linhas):
    mae = lambda k, L: st.mean(abs(l[k]) for l in L)
    sem = [l for l in linhas if not l["flag"]]
    return {"todas": (mae("d25", linhas), mae("d_own", linhas), mae("d_gob", linhas)),
            "sem_flag": (mae("d25", sem), mae("d_own", sem), mae("d_gob", sem)), "n_sem": len(sem),
            "dentro_1_5": (sum(abs(l["d_own"]) <= 1.5 for l in linhas), sum(abs(l["d_gob"]) <= 1.5 for l in linhas))}


def docx(linhas, saida):
    doc = Document()
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21.0), Cm(29.7)
    s.left_margin = s.right_margin = Cm(2.0)
    stl = doc.styles["Normal"]; stl.font.name = "Times New Roman"; stl.font.size = Pt(9)
    stl.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

    def borda(c, **kw):
        tcPr = c._tc.get_or_add_tcPr()
        b = tcPr.find(qn("w:tcBorders"))
        if b is None:
            b = OxmlElement("w:tcBorders"); tcPr.append(b)
        for lado, (v, sz) in kw.items():
            e = OxmlElement(f"w:{lado}"); e.set(qn("w:val"), v); e.set(qn("w:sz"), str(sz)); b.append(e)

    def sombra(c):
        e = OxmlElement("w:shd"); e.set(qn("w:val"), "clear"); e.set(qn("w:color"), "auto"); e.set(qn("w:fill"), "E7E6E6")
        c._tc.get_or_add_tcPr().append(e)

    def cel(c, txt, bold=False, align="r", sub=None):
        c.text = ""
        p = c.paragraphs[0]
        r = p.add_run(txt); r.bold = bold; r.font.size = Pt(9)
        p.alignment = {"l": WD_ALIGN_PARAGRAPH.LEFT, "r": WD_ALIGN_PARAGRAPH.RIGHT, "c": WD_ALIGN_PARAGRAPH.CENTER}[align]
        p.paragraph_format.space_after = Pt(0); p.paragraph_format.space_before = Pt(1)
        if sub:
            p2 = c.add_paragraph(); r2 = p2.add_run(sub); r2.font.size = Pt(7.5)
            p2.alignment = p.alignment; p2.paragraph_format.space_after = Pt(0)

    p = doc.add_paragraph()
    r = p.add_run("Tabela 7 – "); r.bold = True; r.font.size = Pt(11)
    r = p.add_run("Validação externa das projeções estaduais: nosso estudo e projeção do IPEA, base 2025 e 2033"); r.font.size = Pt(11)
    p.paragraph_format.space_after = Pt(6)

    larg = [1.2, 1.5, 1.5, 1.4, 1.5, 1.9, 1.5, 1.5, 1.9]
    t = doc.add_table(rows=0, cols=len(larg)); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); t._tbl.tblPr.append(lay)

    def nova():
        r = t.add_row()
        for i, w in enumerate(larg): r.cells[i].width = Cm(w)
        return r.cells

    c1, c2 = nova(), nova()
    g1 = c1[1].merge(c1[3]); cel(g1, "Base 2025", bold=True, align="c", sub="R$ bilhões"); borda(g1, bottom=("single", 4))
    g2 = c1[4].merge(c1[8]); cel(g2, "Projeção 2033", bold=True, align="c", sub="R$ bilhões"); borda(g2, bottom=("single", 4))
    m = c1[0].merge(c2[0]); cel(m, "UF", bold=True, align="l")
    for j, h in enumerate(["Nosso", "IPEA", "Δ (%)", "Nosso", "Nosso com\ncoef. IPEA", "IPEA", "Δ (%)\nnosso", "Δ (%)\ncoef. IPEA"]):
        cel(c2[1 + j], h, bold=True)
    for c in c1: borda(c, top=("single", 12))
    for c in c2: borda(c, bottom=("single", 6))

    for l in linhas:
        c = nova(); neg = l["uf"] == DESTAQUE
        cel(c[0], l["uf"] + ("¹" if l["flag"] else ""), bold=neg, align="l")
        for k, v in enumerate([br(l["pre"]), br(l["ipea25"]), sg(l["d25"]), br(l["own"]), br(l["gob"]), br(l["ipea33"]), sg(l["d_own"]), sg(l["d_gob"])]):
            cel(c[1 + k], v, bold=neg)
        if neg:
            for x in c: sombra(x)

    rs = resumo(linhas)
    for k, (rot, key) in enumerate([("Diferença média absoluta (27 UFs)", "todas"), (f"Idem, sem as UFs com ¹ ({rs['n_sem']} UFs)", "sem_flag")]):
        c = nova()
        m = c[0].merge(c[2]); cel(m, rot, bold=True, align="l")
        a, b, g = rs[key]
        cel(c[3], sg(a).replace("+", ""), bold=True)
        for x in c[4:7]: cel(x, "")
        cel(c[7], sg(b).replace("+", ""), bold=True); cel(c[8], sg(g).replace("+", ""), bold=True)
        for x in c:
            sombra(x)
            if k == 0: borda(x, top=("single", 6))
            if k == 1: borda(x, bottom=("single", 12))

    n = doc.add_paragraph(); n.paragraph_format.space_before = Pt(4)
    r = n.add_run("Nota: Δ = (IPEA / nosso) − 1. Valores em R$ bilhões, preços de 2025. \"Nosso\" usa o coeficiente de destino próprio (POF e Censo); "
                  "\"Nosso com coef. IPEA\" aplica ao mesmo modelo o coeficiente de destino de Gobetti e Monteiro (2023), com o Seguro-Receita "
                  "recalculado. A diferença média é calculada em valor absoluto. ¹ UFs com diferença de 5% ou mais já na base de 2025 (tratamento "
                  "de fundos estaduais e de receita média), de modo que a divergência de 2033 reflete também a base, e não o coeficiente. "
                  "Ordenada pela diferença na base 2025.")
    r.font.size = Pt(8)
    n2 = doc.add_paragraph()
    r = n2.add_run("Fonte: elaboração própria, com base em SICONFI/STN e nas projeções do IPEA/Gobetti (simulações do estudo para o Comsefaz, 2026).")
    r.font.size = Pt(8)
    doc.save(saida)


if __name__ == "__main__":
    ipea = le_ipea(sys.argv[1])
    linhas = monta(ipea)
    saida = sys.argv[2] if len(sys.argv) > 2 else "tabela-validacao-externa-pr.docx"
    docx(linhas, saida)
    rs = resumo(linhas)
    pr = next(l for l in linhas if l["uf"] == DESTAQUE)
    print({k: (tuple(round(x, 2) for x in v) if isinstance(v, tuple) else v) for k, v in rs.items()})
    print("PR", {k: round(v, 2) for k, v in pr.items() if k != "uf" and k != "flag"}, [l["uf"] for l in linhas if l["flag"]])
