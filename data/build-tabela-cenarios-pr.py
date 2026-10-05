#!/usr/bin/env python3
"""
Tabela do artigo: projeção da receita tributária estadual (ICMS/IBS) do Paraná,
2029-2033, em três cenários para o coeficiente de destino (formato da Tabela 6
do estudo do ES). Um único parâmetro varia: a fonte do coeficiente de destino.

  Própria  : composto robusto de POF, Censo e PNAD (phi-dest-pof-censo.json, phi_fam_pct)
  Neutro   : destino = participação de origem em 2025 (referência, não previsão)
  Gobetti  : Gobetti e Monteiro (2023, IPEA), Tabela 1 (gobetti_tabela1_2023_pct)

Mesmo quadro do Estudo 15 (resultados-consolidados-ibs.json): o contrafactual
de cada ano cresce com o bolo nacional; a variação é contra ele. Fórmula do ente
estadual: R = resid*phi_neutro + (1-ca)*IBS_hist*phi_CPT + IBS_destino_liq*phi_dest
(PR não recebe Seguro-Receita na esfera estadual). phi_CPT é obtido da projeção
publicada (POF/Censo, 2033) e validado contra todos os anos publicados.

Uso: python3 build-tabela-cenarios-pr.py -> auditoria-pr-validacao/tabelas-artigo/tabela-cenarios-receita-estadual-pr.docx
"""
import json
from pathlib import Path
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm
import rateio_consumo_compras as rcc

HERE = Path(__file__).parent
ROOT = HERE.parent
OUT = HERE / "auditoria-pr-validacao" / "tabelas-artigo"
ANOS = [2029, 2030, 2031, 2032, 2033]

R = json.load(open(ROOT / "painelufir/data/resultados-consolidados-ibs.json"))["por_uf_estado"]["PR"]
F = json.load(open(HERE / "phi-dest-pof-censo.json"))
frac_estado = F["frac_estado_pct"] / 100
phi_uf = F["por_uf"]["PR"]
N = {r["ano"]: r for r in json.load(open(ROOT / "painelufir/data/ibs-projecao-nacional.json"))["projecao"]}
ref = {a: N[a]["icms_iss_residual"] + N[a]["ibs_bruto"] for a in N}
c_neutro = R["contrafactual_por_ano"]["2033"] / ref[2033]
c_own, _ = rcc.esferas_uf(phi_uf, False, frac_estado, 0)
c_gob = phi_uf["gobetti_tabela1_2023_pct"] / 100 * frac_estado
n33 = N[2033]
c_cpt = (R["pos_por_ano"]["2033"] - n33["ibs_destino_liquido"] * c_own) / ((1 - n33["ca"]) * n33["ibs_historico"])


def receita(a, c_dest):
    n = N[a]
    return n["icms_iss_residual"] * c_neutro + (1 - n["ca"]) * n["ibs_historico"] * c_cpt + n["ibs_destino_liquido"] * c_dest


for a in ANOS:   # validação: reproduz a projeção publicada (POF/Censo) em todos os anos
    assert abs(receita(a, c_own) - R["pos_por_ano"][str(a)]) < 1e3, a

CEN = [("Estimativa própria", "Composto POF/Censo/PNAD", c_own), ("Neutro", "destino = origem", c_neutro),
       ("Gobetti e Monteiro", "IPEA, 2023", c_gob)]
dados = {a: {"ref": ref[a] * c_neutro, "cen": [receita(a, c) for _, _, c in CEN]} for a in ANOS}


def num(x, casas=2):
    return f"{x:,.{casas}f}".replace(",", "X").replace(".", ",").replace("X", ".")


def sgn(x, casas=1, suf=""):
    return f"{x:+,.{casas}f}{suf}".replace(",", "X").replace(".", ",").replace("X", ".").replace("-", "−")


doc = Document()
s = doc.sections[0]
s.page_width, s.page_height = Cm(21.0), Cm(29.7)
s.left_margin = s.right_margin = Cm(2.0)
st = doc.styles["Normal"]; st.font.name = "Times New Roman"; st.font.size = Pt(9)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")


def borda(c, **kw):
    tcPr = c._tc.get_or_add_tcPr()
    b = tcPr.find(qn("w:tcBorders"))
    if b is None:
        b = OxmlElement("w:tcBorders"); tcPr.append(b)
    for lado, (v, sz) in kw.items():
        e = OxmlElement(f"w:{lado}"); e.set(qn("w:val"), v); e.set(qn("w:sz"), str(sz)); b.append(e)


def cel(c, txt, bold=False, align="r", size=9, sub=None):
    c.text = ""
    p = c.paragraphs[0]
    r = p.add_run(txt); r.bold = bold; r.font.size = Pt(size)
    p.alignment = {"l": WD_ALIGN_PARAGRAPH.LEFT, "r": WD_ALIGN_PARAGRAPH.RIGHT, "c": WD_ALIGN_PARAGRAPH.CENTER}[align]
    p.paragraph_format.space_after = Pt(0); p.paragraph_format.space_before = Pt(1)
    if sub:
        p2 = c.add_paragraph(); r2 = p2.add_run(sub); r2.font.size = Pt(7.5)
        p2.alignment = p.alignment; p2.paragraph_format.space_after = Pt(0)


p = doc.add_paragraph()
r = p.add_run("Tabela 6 – "); r.bold = True; r.font.size = Pt(11)
r = p.add_run("Projeção da receita tributária estadual (ICMS/IBS) do Paraná, 2029–2033"); r.font.size = Pt(11)
p.paragraph_format.space_after = Pt(6)

larg = [1.1, 1.5] + [1.5, 1.35, 1.6] * 3
t = doc.add_table(rows=0, cols=len(larg)); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); t._tbl.tblPr.append(lay)


def nova():
    r = t.add_row()
    for i, w in enumerate(larg): r.cells[i].width = Cm(w)
    return r.cells


c1, c2 = nova(), nova()
for k, (nome, sub, _) in enumerate(CEN):
    g = c1[2 + 3 * k].merge(c1[4 + 3 * k]); cel(g, nome, bold=True, align="c", sub=sub)
    borda(g, bottom=("single", 4))
    for j, h in enumerate(["R$ bi", "Δ%", "Δ R$ mi"]):
        cel(c2[2 + 3 * k + j], h, bold=True, align="r")
for i, h in ((0, "Ano"), (1, "Ref.")):
    m = c1[i].merge(c2[i]); cel(m, h, bold=True, align="l" if i == 0 else "r")
for c in c1: borda(c, top=("single", 12))
for c in c2: borda(c, bottom=("single", 6))

acum = [0.0] * len(CEN)
for a in ANOS:
    c = nova()
    cel(c[0], str(a), align="l"); cel(c[1], num(dados[a]["ref"] / 1e9))
    for k, v in enumerate(dados[a]["cen"]):
        d = v - dados[a]["ref"]; acum[k] += d
        cel(c[2 + 3 * k], num(v / 1e9)); cel(c[3 + 3 * k], sgn(d / dados[a]["ref"] * 100, 1)); cel(c[4 + 3 * k], sgn(d / 1e6, 0))
c = nova()
cel(c[0], "Acum.", bold=True, align="l"); cel(c[1], "")
for k, d in enumerate(acum):
    cel(c[2 + 3 * k], sgn(d / 1e9, 2), bold=True); cel(c[3 + 3 * k], ""); cel(c[4 + 3 * k], sgn(d / 1e6, 0), bold=True)
for x in c:
    borda(x, top=("single", 6), bottom=("single", 12))

n = doc.add_paragraph(); n.paragraph_format.space_before = Pt(4)
r = n.add_run("Nota: Ref. é a receita estadual projetada sem a reforma (participação neutra de 2025 do ente estadual, ICMS + FECOP, líquido de "
              "cota-parte, aplicada à receita de referência nacional de cada ano, em R$ bilhões de 2025). Δ% e Δ R$ mi comparam a receita projetada "
              "com a referência do mesmo ano; Acum. soma os cinco anos. Os cenários diferem apenas na fonte do coeficiente de destino: estimativa "
              "própria (composto de POF, Censo e PNAD Contínua); neutro (destino igual à participação de origem, referência e não previsão; o histórico de transição "
              "continua valendo); e Gobetti e Monteiro (2023, Tabela 1). O Paraná não recebe repasse do Seguro-Receita na esfera estadual.")
r.italic = False; r.font.size = Pt(8)
n2 = doc.add_paragraph()
r = n2.add_run("Fonte: SICONFI/STN; EC n.º 132/2023; LC n.º 214/2025; LC n.º 227/2026; Gobetti e Monteiro (2023). Elaboração própria.")
r.font.size = Pt(8)
doc.save(OUT / "tabela-cenarios-receita-estadual-pr.docx")

print("coef estado %: própria", round(c_own * 100, 3), "neutro", round(c_neutro * 100, 3), "gobetti", round(c_gob * 100, 3))
for a in ANOS:
    print(a, num(dados[a]["ref"] / 1e9), [(num(v / 1e9), sgn((v / dados[a]["ref"] - 1) * 100, 1), sgn((v - dados[a]["ref"]) / 1e6, 0)) for v in dados[a]["cen"]])
print("acum mi", [round(x / 1e6) for x in acum])
