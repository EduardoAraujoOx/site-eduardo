#!/usr/bin/env python3
"""
Tabela do artigo: receita estadual projetada por UF no cenário de destino pleno,
2029-2033 (R$ bilhões, preços de 2025), com o Paraná destacado. Lê
painelufir/data/resultados-consolidados-ibs.json (por_uf_estado, Estudo 15).

Uso: python3 build-tabela-uf.py -> auditoria-pr-validacao/tabelas-artigo/tabela-receita-estadual-uf-2029-2033.docx
"""
import json
from pathlib import Path
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm

HERE = Path(__file__).parent
SRC = HERE.parent / "painelufir" / "data" / "resultados-consolidados-ibs.json"
OUT = HERE / "auditoria-pr-validacao" / "tabelas-artigo"
DESTAQUE = "PR"


def num(x, casas=1):
    return f"{x:,.{casas}f}".replace(",", "X").replace(".", ",").replace("X", ".")


def pct(v):
    v = round(v * 100, 1)
    return ("+0,0%" if v == 0 else (f"{v:+.1f}%".replace(".", ",").replace("-", "−")))


d = json.load(open(SRC))
anos = d["anos"]
linhas = sorted(d["por_uf_estado"].items(), key=lambda kv: kv[1]["variacao_por_ano"][str(anos[-1])])

doc = Document()
s = doc.sections[0]
s.page_width, s.page_height = Cm(21.0), Cm(29.7)
s.left_margin = s.right_margin = Cm(1.95)
st = doc.styles["Normal"]; st.font.name = "Times New Roman"; st.font.size = Pt(9)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")


def borda(c, **kw):
    tcPr = c._tc.get_or_add_tcPr()
    b = tcPr.find(qn("w:tcBorders"))
    if b is None:
        b = OxmlElement("w:tcBorders"); tcPr.append(b)
    for lado, (v, sz) in kw.items():
        e = OxmlElement(f"w:{lado}"); e.set(qn("w:val"), v); e.set(qn("w:sz"), str(sz)); b.append(e)


def sombra(c, cor="E7E6E6"):
    e = OxmlElement("w:shd"); e.set(qn("w:val"), "clear"); e.set(qn("w:color"), "auto"); e.set(qn("w:fill"), cor)
    c._tc.get_or_add_tcPr().append(e)


def cel(c, txt, bold=False, align="r", size=9, sub=None):
    c.text = ""
    p = c.paragraphs[0]
    r = p.add_run(txt); r.bold = bold; r.font.size = Pt(size)
    p.alignment = {"l": WD_ALIGN_PARAGRAPH.LEFT, "r": WD_ALIGN_PARAGRAPH.RIGHT, "c": WD_ALIGN_PARAGRAPH.CENTER}[align]
    p.paragraph_format.space_after = Pt(0); p.paragraph_format.space_before = Pt(1)
    if sub:
        p2 = c.add_paragraph(); r2 = p2.add_run(sub); r2.font.size = Pt(7.5)
        p2.alignment = p.alignment; p2.paragraph_format.space_after = Pt(0)


p = doc.add_paragraph()
r = p.add_run("Tabela 5 – "); r.bold = True; r.font.size = Pt(11)
r = p.add_run("Receita estadual projetada por UF no cenário de destino pleno, 2029–2033"); r.font.size = Pt(11)
p.paragraph_format.space_after = Pt(6)

larg = [1.0, 1.7] + [1.3] * 5 + [1.55] * 5
t = doc.add_table(rows=0, cols=len(larg)); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); t._tbl.tblPr.append(lay)


def nova():
    r = t.add_row()
    for i, w in enumerate(larg): r.cells[i].width = Cm(w)
    return r.cells


# cabeçalho: dois níveis
c1, c2 = nova(), nova()
g1 = c1[2].merge(c1[6]); cel(g1, "Receita pós-IBS", bold=True, align="c", sub="R$ bilhões, preços de 2025")
g2 = c1[7].merge(c1[11]); cel(g2, "Variação da receita", bold=True, align="c", sub="ano contra pré-IBS")
borda(g1, bottom=("single", 4)); borda(g2, bottom=("single", 4))
for i, h in ((0, "UF"), (1, "Receita\npré-IBS")):
    m = c1[i].merge(c2[i]); cel(m, h, bold=True, align="l" if i == 0 else "r")
for k, a in enumerate(anos):
    cel(c2[2 + k], str(a), bold=True, align="r"); cel(c2[7 + k], str(a), bold=True, align="r")
for c in c1: borda(c, top=("single", 12))
for c in c2: borda(c, bottom=("single", 6))

for i, (uf, v) in enumerate(linhas):
    c = nova()
    neg = uf == DESTAQUE
    cel(c[0], uf, bold=neg, align="l"); cel(c[1], num(v["pre_2025"] / 1e9), bold=neg)
    for k, a in enumerate(anos):
        cel(c[2 + k], num(v["pos_por_ano"][str(a)] / 1e9), bold=neg)
        cel(c[7 + k], pct(v["variacao_por_ano"][str(a)]), bold=neg)
    if neg:
        for x in c: sombra(x)
    if i == len(linhas) - 1:
        for x in c: borda(x, bottom=("single", 12))

n = doc.add_paragraph()
n.paragraph_format.space_before = Pt(4)
r = n.add_run("Fonte: SICONFI/STN; EC n.º 132/2023; LC n.º 214/2025; LC n.º 227/2026. Elaboração própria. "
              "Nota: receita pré-IBS corresponde à participação neutra de 2025 do ente estadual (ICMS + FECOP, líquido de cota-parte) "
              "aplicada à receita de referência nacional de cada ano. Ordenada pela variação relativa em 2033, da mais negativa à mais "
              "positiva. Paraná em destaque.")
r.font.size = Pt(8)
doc.save(OUT / "tabela-receita-estadual-uf-2029-2033.docx")
pr = dict(linhas)[DESTAQUE]
print(len(linhas), DESTAQUE, num(pr["pre_2025"] / 1e9), [num(pr["pos_por_ano"][str(a)] / 1e9) for a in anos],
      [pct(pr["variacao_por_ano"][str(a)]) for a in anos], [u for u, _ in linhas].index(DESTAQUE) + 1)
