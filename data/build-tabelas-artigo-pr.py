#!/usr/bin/env python3
"""
Tabelas de dispersão municipal per capita do Paraná para o artigo (Word).

Mesma técnica de dispersao_intramunicipal (build-resultados-consolidados.py):
mostra os extremos (municípios de maior e menor valor per capita em 2025) e
projeta os mesmos municípios em 2033 e 2077, para verificar se a distância
entre eles diminui conforme o critério de destino substitui o histórico.

Diferenças em relação ao Estudo 15 (nacional): usa a base municipal final
auditada do Paraná (data/auditoria-pr-validacao/municipios-pr-final-auditado-
latest.csv) e o ISS ajustado ano a ano (data/auditoria-dca-pr/observacoes-pr-
latest.csv) em vez das bases nacionais canônicas.

Fórmula de 2033 e 2077 (idêntica a project_municipio, com ibs-projecao-longo-
prazo.json): receita = (1-ca)*IBS_histórico*CPT + IBS_destino_líquido*coef_destino
+ (1-ca)*Seguro-Receita. Em 2033 e 2077 o resíduo ICMS/ISS é zero (sa=1).
Seguro-Receita 2077 por município: fatia do município no fundo municipal do PR
em 2033 (base auditada) x fundo municipal do PR em 2077 (seguro-receita-
repasses-longo-prazo.json), mesma aproximação do Estudo 15. R$ constantes de
2025; população fixada na média 2019-2026 nos três pontos.

Uso: python3 build-tabelas-artigo-pr.py
"""

import csv
import json
import statistics
from pathlib import Path

HERE = Path(__file__).parent
ROOT = HERE.parent
OUT = HERE / "auditoria-pr-validacao" / "tabelas-artigo"
OUT.mkdir(parents=True, exist_ok=True)

FINAL = HERE / "auditoria-pr-validacao" / "municipios-pr-final-auditado-latest.csv"
OBS = HERE / "auditoria-dca-pr" / "observacoes-pr-latest.csv"
PAINEL = ROOT / "painelufir" / "data" / "painel-municipios" / "PR.json"
LP = ROOT / "painelufir" / "data" / "ibs-projecao-longo-prazo.json"
SEG_LP = ROOT / "painelufir" / "data" / "seguro-receita-repasses-longo-prazo.json"

N_EXTREMOS = 10
SALTO_MAX = 10.0


def carregar():
    final = {r["codigo_ibge"]: r for r in csv.DictReader(open(FINAL, encoding="utf-8-sig"))}
    iss25, hist = {}, {}
    for r in csv.DictReader(open(OBS, encoding="utf-8-sig")):
        if r["adjusted_iss"] in ("", None):
            continue
        if r["ano"] == "2025":
            iss25[r["codigo_ibge"]] = float(r["adjusted_iss"])
        else:
            hist.setdefault(r["codigo_ibge"], []).append(float(r["adjusted_iss"]))
    # razão ISS 2025 / mediana 2019-2024 (triagem de salto atípico)
    salto = {c: iss25[c] / statistics.median(v) for c, v in hist.items()
             if c in iss25 and len(v) >= 3 and statistics.median(v) > 0}
    painel = json.load(open(PAINEL))["municipios"]
    lp = {p["ano"]: p for p in json.load(open(LP))["projecao"]}
    seg = json.load(open(SEG_LP))["anos"]
    return final, iss25, salto, painel, lp, seg


def montar():
    final, iss25, salto, painel, lp, seg = carregar()
    fundo_pr_2077 = seg["2077"]["repasse_por_uf"]["PR"]["municipio"]
    soma_seg_2033 = sum(float(r["seguro_2033_final"]) for r in final.values())
    p33, p77 = lp[2033], lp[2077]

    linhas = []
    for cod, r in final.items():
        pop = (painel.get(cod) or {}).get("pop_media") or 0
        if not pop:
            continue
        cpt = float(r["cpt_final_pct"]) / 100
        dest = float(r["destino_pct"]) / 100
        seg33 = float(r["seguro_2033_final"])
        seg77 = (seg33 / soma_seg_2033) * fundo_pr_2077 if soma_seg_2033 else 0.0
        rec33 = float(r["receita_2033_final"])
        rec77 = ((1 - p77["ca"]) * p77["ibs_historico"] * cpt
                 + p77["ibs_destino_liquido"] * dest + (1 - p77["ca"]) * seg77)
        # confere a reprodução de 2033 com o CSV auditado
        rec33_chk = ((1 - p33["ca"]) * p33["ibs_historico"] * cpt
                     + p33["ibs_destino_liquido"] * dest + (1 - p33["ca"]) * seg33)
        assert abs(rec33_chk - rec33) < 1.0, (cod, rec33_chk, rec33)
        linhas.append({
            "codigo_ibge": cod, "municipio": r["municipio"], "pop": pop,
            "iss_2025": iss25.get(cod), "base_2025": float(r["base_2025_final"]),
            "rec_2033": rec33, "rec_2077": rec77,
            "iss_direto": r["iss_2025_direto"] == "True",
            "pendente": bool(r["iss_pendente_anos"]),
            "obs_anos": int(r["iss_observado_anos"] or 0),
            "amostra_artigo": r["qualificado_artigo"] == "True",
            "salto_iss": salto.get(cod),
        })
    for l in linhas:
        l["iss_pc"] = l["iss_2025"] / l["pop"] if l["iss_2025"] is not None else None
        l["pc_2025"] = l["base_2025"] / l["pop"]
        l["pc_2033"] = l["rec_2033"] / l["pop"]
        l["pc_2077"] = l["rec_2077"] / l["pop"]
    # elegíveis: ISS 2025 direto, sem pendência de auditoria, >=5 anos observados
    # e sem salto atípico (ISS 2025 > 10x a mediana 2019-2024: hoje só
    # Guaraqueçaba, de R$ 0,5 mi para R$ 23,8 mi -- provável erro de declaração,
    # a confirmar no PIT/TCE-PR; fica fora dos extremos até essa confirmação).
    elegiveis = [l for l in linhas if l["iss_direto"] and not l["pendente"]
                 and l["obs_anos"] >= 5 and l["iss_pc"] is not None
                 and (l["salto_iss"] or 0) <= SALTO_MAX]
    return linhas, elegiveis


def gini(v):
    v = sorted(v)
    n, s = len(v), sum(v)
    return (2 * sum((i + 1) * x for i, x in enumerate(v)) / (n * s)) - (n + 1) / n if s else 0.0


def pctil(v, q):
    v = sorted(v)
    k = (len(v) - 1) * q
    f = int(k)
    return v[f] + (v[min(f + 1, len(v) - 1)] - v[f]) * (k - f)


def dispersao(base, campo):
    v = [l[campo] for l in base]
    return {
        "max": max(v), "min": min(v), "amplitude": max(v) - min(v),
        "max_min": max(v) / min(v), "p90_p10": pctil(v, .9) / pctil(v, .1),
        "cv": statistics.pstdev(v) / statistics.mean(v), "gini": gini(v),
    }


def resumo_dispersao(base):
    return {c: dispersao(base, c) for c in ("pc_2025", "pc_2033", "pc_2077")}


def main():
    linhas, elegiveis = montar()
    amostra = [l for l in elegiveis if l["amostra_artigo"]]

    por_iss = sorted(elegiveis, key=lambda l: -l["iss_pc"])
    por_rec = sorted(elegiveis, key=lambda l: -l["pc_2025"])
    tab1 = por_iss[:N_EXTREMOS] + por_iss[-N_EXTREMOS:]
    tab2 = por_rec[:N_EXTREMOS] + por_rec[-N_EXTREMOS:]

    def razao(l, a, b):
        return l[b] / l[a]

    resumo = {
        "n_municipios_pr": len(linhas), "n_elegiveis": len(elegiveis),
        "n_amostra_artigo": len(amostra),
        "iss_pc_2025": {
            "max": por_iss[0]["iss_pc"], "min": por_iss[-1]["iss_pc"],
            "max_min": por_iss[0]["iss_pc"] / por_iss[-1]["iss_pc"],
            "mediana": statistics.median(l["iss_pc"] for l in elegiveis),
            "p90_p10": pctil([l["iss_pc"] for l in elegiveis], .9) / pctil([l["iss_pc"] for l in elegiveis], .1),
            "cv": statistics.pstdev([l["iss_pc"] for l in elegiveis]) / statistics.mean(l["iss_pc"] for l in elegiveis),
        },
        "dispersao_receita_pc_elegiveis": resumo_dispersao(elegiveis),
        "dispersao_receita_pc_amostra_artigo": resumo_dispersao(amostra),
        "tab1_iss": [{k: l[k] for k in ("municipio", "pop", "iss_pc", "pc_2025", "pc_2033", "pc_2077")} for l in tab1],
        "tab2_receita": [{k: l[k] for k in ("municipio", "pop", "iss_pc", "pc_2025", "pc_2033", "pc_2077")} for l in tab2],
    }
    json.dump(resumo, open(OUT / "resumo-tabelas.json", "w"), ensure_ascii=False, indent=1)

    with open(OUT / "municipios-percapita-pr.csv", "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["codigo_ibge", "municipio", "pop_media", "iss_2025_rs", "iss_per_capita_2025",
                    "receita_per_capita_2025", "receita_per_capita_2033", "receita_per_capita_2077",
                    "elegivel_ranking", "amostra_artigo_188"])
        elegiveis_ids = {l["codigo_ibge"] for l in elegiveis}
        for l in sorted(linhas, key=lambda l: -(l["iss_pc"] or -1)):
            w.writerow([l["codigo_ibge"], l["municipio"], round(l["pop"], 1),
                        round(l["iss_2025"], 2) if l["iss_2025"] is not None else "",
                        round(l["iss_pc"], 2) if l["iss_pc"] is not None else "",
                        round(l["pc_2025"], 2), round(l["pc_2033"], 2), round(l["pc_2077"], 2),
                        l["codigo_ibge"] in elegiveis_ids, l["amostra_artigo"]])
    return resumo, tab1, tab2


# ── Word ────────────────────────────────────────────────────────────────────
def _br(x, casas=0):
    return f"{x:,.{casas}f}".replace(",", "X").replace(".", ",").replace("X", ".")


def gerar_docx(resumo):
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, Cm

    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Cm(2.5)
    st = doc.styles["Normal"]
    st.font.name, st.font.size = "Times New Roman", Pt(11)

    def par(txt, bold=False, italic=False, size=None, align=None, after=4):
        p = doc.add_paragraph()
        r = p.add_run(txt)
        r.bold, r.italic = bold, italic
        if size:
            r.font.size = Pt(size)
        if align:
            p.alignment = align
        p.paragraph_format.space_after = Pt(after)
        return p

    def borda(cell, **kw):
        tcPr = cell._tc.get_or_add_tcPr()
        b = OxmlElement("w:tcBorders")
        for lado, v in kw.items():
            e = OxmlElement(f"w:{lado}")
            e.set(qn("w:val"), v)
            e.set(qn("w:sz"), "6")
            b.append(e)
        tcPr.append(b)

    def tabela(cab, linhas, larguras, sep_apos=None, negrito_ultima=False):
        t = doc.add_table(rows=1, cols=len(cab))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, h in enumerate(cab):
            c = t.rows[0].cells[i]
            c.text = ""
            r = c.paragraphs[0].add_run(h)
            r.bold, r.font.size = True, Pt(9)
            c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER if i else WD_ALIGN_PARAGRAPH.LEFT
            borda(c, top="single", bottom="single")
        for n, ln in enumerate(linhas):
            cells = t.add_row().cells
            for i, v in enumerate(ln):
                cells[i].text = ""
                r = cells[i].paragraphs[0].add_run(str(v))
                r.font.size = Pt(9)
                cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT if i else WD_ALIGN_PARAGRAPH.LEFT
                if sep_apos is not None and n == sep_apos:
                    borda(cells[i], bottom="dashed")
                if n == len(linhas) - 1:
                    borda(cells[i], bottom="single")
        for row in t.rows:
            for i, w in enumerate(larguras):
                row.cells[i].width = Cm(w)
            for c in row.cells:
                for p in c.paragraphs:
                    p.paragraph_format.space_after = Pt(0)
        return t

    def bloco_extremos(chave):
        d = resumo[chave]
        n = N_EXTREMOS
        linhas = []
        for i, l in enumerate(d):
            linhas.append([f"{i + 1 if i < n else resumo['n_elegiveis'] - 2 * n + i + 1}. {l['municipio']}", _br(l["pop"]),
                           _br(l["iss_pc"]), _br(l["pc_2025"]), _br(l["pc_2033"]), _br(l["pc_2077"])])
        return linhas

    cab = ["Município", "População\n(média 2019-26)", "ISS per capita\n2025", "Receita per capita\n2025", "2033", "2077"]
    larg = [5.2, 2.4, 2.4, 2.6, 1.7, 1.7]

    par("Dispersão municipal per capita no Paraná: 2025, 2033 e 2077", bold=True, size=13, after=8)

    par("Tabela 1 – Municípios paranaenses de maior e menor ISS per capita em 2025 e a receita per capita "
        "(ISS + cota-parte do ICMS, substituídos pelo IBS) projetada em 2033 e 2077 (R$ de 2025)", bold=True, size=10)
    tabela(cab, bloco_extremos("tab1_iss"), larg, sep_apos=N_EXTREMOS - 1)
    par("Fonte: elaboração própria a partir da DCA/Siconfi, Portal da Transparência do PR e PIT/TCE-PR (base municipal final auditada). "
        "Ranking pelo ISS per capita de 2025; a receita per capita de 2033 e 2077 refere-se ao conjunto ISS + cota-parte do ICMS "
        "(o ISS deixa de existir isoladamente em 2033).", italic=True, size=8, after=10)

    par("Tabela 2 – Municípios paranaenses de maior e menor receita per capita (ISS + cota-parte do ICMS) em 2025 e sua "
        "projeção para 2033 e 2077 (R$ de 2025)", bold=True, size=10)
    tabela(cab, bloco_extremos("tab2_receita"), larg, sep_apos=N_EXTREMOS - 1)
    par("Fonte: elaboração própria (mesma base da Tabela 1). População fixada na média 2019-2026 nos três anos, "
        "para isolar o efeito da redistribuição do efeito do crescimento demográfico.", italic=True, size=8, after=10)

    par("Tabela 3 – Medidas de dispersão da receita per capita municipal, Paraná, 2025, 2033 e 2077", bold=True, size=10)
    medidas = [("Maior valor (R$)", "max", 0), ("Menor valor (R$)", "min", 0), ("Amplitude (R$)", "amplitude", 0),
               ("Razão máx./mín.", "max_min", 1), ("Razão P90/P10", "p90_p10", 2),
               ("Coeficiente de variação", "cv", 2), ("Índice de Gini", "gini", 3)]
    for rot, chave in ((f"Painel A – {resumo['n_elegiveis']} municípios com dados auditados", "dispersao_receita_pc_elegiveis"),
                       (f"Painel B – amostra do artigo ({resumo['n_amostra_artigo']} municípios, população média ≥ 10 mil)",
                        "dispersao_receita_pc_amostra_artigo")):
        par(rot, bold=True, size=9, after=2)
        d = resumo[chave]
        linhas = [[nome] + [_br(d[c][k], casas) for c in ("pc_2025", "pc_2033", "pc_2077")] for nome, k, casas in medidas]
        tabela(["Medida", "2025", "2033", "2077"], linhas, [6.0, 3.0, 3.0, 3.0])
        par("", after=4)
    par("Fonte: elaboração própria. Valores em R$ constantes de 2025; população fixa (média 2019-2026); 2033 e 2077 conforme o "
        "cronograma do art. 131 do ADCT (EC 132/2023) e crescimento real do PIB de 2,2% a.a. após 2033. Excluídos: municípios sem "
        "ISS de 2025 observado diretamente, com pendência de auditoria, com menos de 5 anos de ISS observado ou com salto atípico "
        "do ISS em 2025 (Guaraqueçaba, a confirmar no PIT/TCE-PR).", italic=True, size=8)

    doc.save(OUT / "tabelas-dispersao-percapita-pr.docx")


if __name__ == "__main__":
    r, t1, t2 = main()
    gerar_docx(r)
    print(json.dumps({k: r[k] for k in ("n_municipios_pr", "n_elegiveis", "n_amostra_artigo", "iss_pc_2025")}, indent=1, ensure_ascii=False))
    for nome in ("dispersao_receita_pc_elegiveis", "dispersao_receita_pc_amostra_artigo"):
        print(nome)
        for c, d in r[nome].items():
            print(" ", c, {k: round(v, 3) for k, v in d.items()})
    for nome, t in (("TAB1 (rank ISS pc)", t1), ("TAB2 (rank receita pc)", t2)):
        print(nome)
        for l in t:
            print(f"  {l['municipio']:<26} pop={l['pop']:>9.0f} iss_pc={l['iss_pc']:>8.0f} 25={l['pc_2025']:>7.0f} 33={l['pc_2033']:>7.0f} 77={l['pc_2077']:>7.0f}")
